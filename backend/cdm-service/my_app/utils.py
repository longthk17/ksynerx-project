from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from rest_framework import status
from rest_framework.response import Response
from rest_framework.exceptions import ValidationError


def response(
    *,
    data=None,
    message='Success',
    errors=None,
    status_code=status.HTTP_200_OK,
):
    if status_code >= 400:
        body = {
            'code': status_code,
            'errorMessage': message,
        }

        if errors is not None:
            body['errors'] = errors
    else:
        body = data

    return Response(data=body, status=status_code)


import arrow


def get_relative_datetime(
    *,
    set_year=None,
    set_month=None,
    set_day=None,
    set_hour=None,
    set_minute=None,
    set_second=None,
    shift_years=0,
    shift_months=0,
    shift_days=0,
    shift_hours=0,
    shift_minutes=0,
    shift_seconds=0,
    timezone="Asia/Ho_Chi_Minh",
):
    """Set ngày/giờ trước, sau đó dịch thời gian."""

    set_values = {
        "year": set_year,
        "month": set_month,
        "day": set_day,
        "hour": set_hour,
        "minute": set_minute,
        "second": set_second,
    }

    set_values = {key: value for key, value in set_values.items() if value is not None}

    target = (
        arrow.now(timezone)
        .replace(**set_values, microsecond=0)
        .shift(
            years=shift_years,
            months=shift_months,
            days=shift_days,
            hours=shift_hours,
            minutes=shift_minutes,
            seconds=shift_seconds,
        )
    )

    return target.datetime


def read_excel_sheets(uploaded_file):
    try:
        workbook = load_workbook(
            uploaded_file,
            read_only=True,
            data_only=True,
        )
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise ValidationError({"file": f"Error processing Excel file: {exc}"}) from exc

    try:
        sheets = {}

        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header_row = next(rows, None)

            if header_row is None:
                sheets[sheet.title] = []
                continue

            headers = [
                str(value).strip() if value is not None else "" for value in header_row
            ]

            named_headers = [header for header in headers if header]

            if len(named_headers) != len(set(named_headers)):
                raise ValidationError(
                    {"file": f"Sheet '{sheet.title}' có tên cột trùng."}
                )

            records = []

            for row_number, row in enumerate(rows, start=2):
                if all(
                    value is None or isinstance(value, str) and not value.strip()
                    for value in row
                ):
                    continue

                records.append(
                    {
                        "row": row_number,
                        "data": {
                            header: value
                            for header, value in zip(headers, row)
                            if header
                        },
                    }
                )

            sheets[sheet.title] = records

        return sheets

    finally:
        workbook.close()


class ExcelMappingError(ValueError):
    def __init__(self, errors):
        self.errors = errors
        super().__init__("Dữ liệu Excel không hợp lệ")


def map_product_sheets(sheets):
    required_sheets = ("products", "categories", "product_units")
    errors = []
    products_by_sku = {}

    def add_error(sheet, row, field, message):
        errors.append(
            {
                "sheet": sheet,
                "row": row,
                "field": field,
                "message": message,
            }
        )

    def get_sku(data, sheet, row):
        sku = data.get("sku")

        if not isinstance(sku, str) or not sku.strip():
            add_error(sheet, row, "sku", "SKU không hợp lệ.")
            return None

        return sku.strip()

    for sheet_name in required_sheets:
        if sheet_name not in sheets:
            add_error(
                sheet_name,
                None,
                None,
                f"Thiếu sheet '{sheet_name}'.",
            )

    if errors:
        raise ExcelMappingError(errors)

    # 1. Tạo sản phẩm theo SKU.
    for record in sheets["products"]:
        data = record["data"].copy()
        row = record["row"]
        sku = get_sku(data, "products", row)

        if sku is None:
            continue

        if sku in products_by_sku:
            add_error(
                "products",
                row,
                "sku",
                f"SKU '{sku}' bị trùng trong sheet products.",
            )
            continue

        data["sku"] = sku
        data["categories"] = []
        data["productUnits"] = []

        products_by_sku[sku] = data

    # 2. Ghép dữ liệu từ hai sheet con.
    mappings = (
        ("categories", "categories", "categoryCode"),
        ("product_units", "productUnits", "unitCode"),
    )

    for sheet_name, target_field, code_field in mappings:
        seen = set()

        for record in sheets[sheet_name]:
            data = record["data"].copy()
            row = record["row"]
            sku = get_sku(data, sheet_name, row)

            if sku is None:
                continue

            if sku not in products_by_sku:
                add_error(
                    sheet_name,
                    row,
                    "sku",
                    f"SKU '{sku}' không có trong sheet products.",
                )
                continue

            code = data.get(code_field)

            if not isinstance(code, str) or not code.strip():
                add_error(
                    sheet_name,
                    row,
                    code_field,
                    f"Field '{code_field}' không hợp lệ.",
                )
                continue

            code = code.strip()
            key = (sku, code)

            if key in seen:
                add_error(
                    sheet_name,
                    row,
                    code_field,
                    f"Mã '{code}' bị trùng trong sản phẩm '{sku}'.",
                )
                continue

            seen.add(key)

            data.pop("sku")
            data[code_field] = code

            products_by_sku[sku][target_field].append(data)

    if not products_by_sku:
        add_error(
            "products",
            None,
            None,
            "Không có sản phẩm để xử lý.",
        )

    if errors:
        raise ExcelMappingError(errors)

    return list(products_by_sku.values())
